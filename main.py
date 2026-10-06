import re
from pathlib import Path
from tkinter import Tk, filedialog

from pypdf import PdfReader
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def select_folder() -> Path | None:
    """Zobrazí dialog pro výběr složky."""
    root = Tk()
    root.withdraw()
    root.attributes("-topmost", True)

    folder = filedialog.askdirectory(
        title="Vyberte složku s váženkami"
    )

    root.destroy()

    if not folder:
        return None

    return Path(folder)


def extract_pdf_text(pdf_path: Path) -> str:
    """Načte text ze všech stran PDF."""
    reader = PdfReader(pdf_path)

    pages = []

    for page in reader.pages:
        text = page.extract_text()

        if text:
            pages.append(text)

    return "\n".join(pages)


def extract_ticket_number(text: str) -> int:
    """
    Najde číslo váženky.

    Např.:
        2026/0031966 -> 31966
        2026\\0031966 -> 31966
    """

    match = re.search(
        r"(?:Váženka\s+a\s+dodací\s+list\s*)?"
        r"\d{4}\s*[\\/]\s*0*(\d+)",
        text,
        re.IGNORECASE,
    )

    if match:
        return int(match.group(1))

    # Záložní varianta pro případ, že extrakce PDF odstraní lomítko.
    match = re.search(
        r"Váženka\s+a\s+dodací\s+list\s+"
        r"\d{4}0*(\d{5,7})",
        text,
        re.IGNORECASE,
    )

    if match:
        return int(match.group(1))

    raise ValueError("Číslo váženky nebylo nalezeno.")


def extract_supplier_code(text: str) -> str:
    """
    Najde kód dodavatele.

    V PDF může být text extrahován například jako:

        Kód dodavatele
        SLOVÁK JAROSLAV
        E201

    Proto nehledáme první slovo za "Kód dodavatele",
    ale kód ve formátu písmeno + číslice.
    """

    match = re.search(
        r"K[oó]d\s+dodavatele[\s\S]{0,100}?\b([A-Z]\d{2,6})\b",
        text,
        re.IGNORECASE,
    )

    if match:
        return match.group(1).upper()

    raise ValueError("Kód dodavatele nebyl nalezen.")


def extract_date(text: str) -> str:
    """
    Najde datum a čas.

    Např.:
        02.10.2026 14:07
    """

    match = re.search(
        r"\b"
        r"(\d{1,2}\.\d{1,2}\.\d{4}"
        r"(?:\s+\d{1,2}:\d{2})?)"
        r"\b",
        text,
    )

    if match:
        return match.group(1)

    raise ValueError("Datum nebylo nalezeno.")


def extract_netto(text: str) -> float:
    """
    Najde Netto v kg a převede jej na tuny.

    Např.:
        Netto [kg] 27480

    Výsledek:
        27.480
    """

    match = re.search(
        r"Netto\s*\[\s*kg\s*\]\s*[:\-]?\s*"
        r"([\d\s]+(?:[,.]\d+)?)",
        text,
        re.IGNORECASE,
    )

    if not match:
        raise ValueError("Netto nebylo nalezeno.")

    value = match.group(1)

    # Odstranění případných mezer mezi tisíci.
    value = value.replace(" ", "")
    value = value.replace(",", ".")

    netto_kg = float(value)

    return netto_kg / 1000.0


def process_pdf(pdf_path: Path) -> dict:
    """Zpracuje jednu váženku."""

    text = extract_pdf_text(pdf_path)

    return {
        "number": extract_ticket_number(text),
        "supplier_code": extract_supplier_code(text),
        "date": extract_date(text),
        "netto": extract_netto(text),
    }


def create_excel(data: list[dict], output_path: Path):
    """Vytvoří výsledný Excel."""

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Váženky"

    # Hlavička
    headers = [
        "Číslo",
        "Kód dodavatele",
        "Datum",
        "Netto [t]",
    ]

    sheet.append(headers)

    # Tučná hlavička
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    # Data
    for item in data:
        sheet.append(
            [
                item["number"],
                item["supplier_code"],
                item["date"],
                item["netto"],
            ]
        )

    # Netto zobrazíme na 3 desetinná místa.
    for row in range(2, sheet.max_row + 1):
        sheet.cell(
            row=row,
            column=4
        ).number_format = "0.000"

    # Automatická šířka sloupců.
    for column in range(1, sheet.max_column + 1):
        max_length = 0

        column_letter = get_column_letter(column)

        for cell in sheet[column_letter]:
            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        sheet.column_dimensions[column_letter].width = max_length + 3

    workbook.save(output_path)


def main():
    print("Import váženek")
    print("===============")
    print()

    folder = select_folder()

    if folder is None:
        print("Nebyla vybrána žádná složka.")
        return

    print(f"Vybraná složka: {folder}")
    print()

    # Najdeme všechny PDF ve vybrané složce.
    pdf_files = sorted(folder.glob("*.pdf"))

    if not pdf_files:
        print(
            "Ve vybrané složce nebyly nalezeny "
            "žádné PDF soubory."
        )
        return

    print(f"Nalezeno PDF souborů: {len(pdf_files)}")
    print()

    data = []
    errors = []

    for pdf_path in pdf_files:
        print(f"Zpracovávám: {pdf_path.name}")

        try:
            item = process_pdf(pdf_path)

            data.append(item)

            print(
                f"  OK: "
                f"{item['number']} | "
                f"{item['supplier_code']} | "
                f"{item['date']} | "
                f"{item['netto']:.3f} t"
            )

        except Exception as e:
            errors.append(
                (
                    pdf_path.name,
                    str(e)
                )
            )

            print(f"  CHYBA: {e}")

    print()

    if not data:
        print("Nepodařilo se zpracovat žádnou váženku.")
        return

    # Seřadíme výsledky podle čísla váženky.
    data.sort(
        key=lambda x: x["number"]
    )

    # Výstupní Excel bude ve stejné složce jako PDF.
    output_path = folder / "vazenky.xlsx"

    create_excel(
        data,
        output_path
    )

    print("Hotovo.")
    print()
    print(f"Excel vytvořen: {output_path}")
    print(f"Úspěšně zpracováno: {len(data)}")
    print(f"Chyby: {len(errors)}")

    if errors:
        print()
        print("Soubory s chybou:")

        for filename, error in errors:
            print(
                f"  {filename}: {error}"
            )


if __name__ == "__main__":
    main()